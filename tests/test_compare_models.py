"""Base vs base + adapter vs gold, per document (testing.compare_models, testing.comparison)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from artifact_registry import paths
from artifact_registry.blob_client import BlobClient, InMemoryBackend
from inference_core.model_runner import EchoBackend, load_model, without_adapter
from testing.comparison import compare, keyed_flatten, write_comparison, write_overview


def _env(value, page=1):
    return {"raw": str(value), "parsed": value, "page_ref": [page]}


GOLD = {
    "policy": {"policy_number": _env("PA-1"), "effective_date": _env("01/01/2026")},
    "auto": {"vehicles": [{"vin": _env("1HGCM82633A004352", 2), "year": _env(2019, 2)},
                          {"vin": _env("5TFDW5F11KX778685", 2), "year": _env(2021, 2)}]},
}


# --------------------------------------------------------------------------
# The comparison
# --------------------------------------------------------------------------


def test_table_rows_are_matched_by_identifier_not_position():
    reordered = {"auto": {"vehicles": [GOLD["auto"]["vehicles"][1], GOLD["auto"]["vehicles"][0]]}}
    flat = keyed_flatten(reordered, GOLD)
    assert "auto.vehicles[vin=1hgcm82633a004352].year" in flat
    assert flat["auto.vehicles[vin=5tfdw5f11kx778685].year"]["parsed"] == 2021


def test_each_model_value_gets_a_result_against_gold():
    base = {"policy": {"policy_number": _env("PA-1"), "program_name": _env("Preferred")}}   # 1 right, 1 invented
    adapter = {"policy": {"policy_number": _env("PA-1"), "effective_date": _env("01/02/2026")},
               "auto": {"vehicles": [{"vin": _env("1HGCM82633A004352"), "year": _env(2019)}]}}
    result = compare(GOLD, base, adapter, lob="personal_auto")
    rows = {row.path: row for row in result.rows}
    assert rows["policy.policy_number"].base_result == "correct"
    assert rows["policy.effective_date"].base_result == "missed"
    assert rows["policy.effective_date"].adapter_result == "wrong"
    assert rows["policy.program_name"].base_result == "invented"
    assert rows["auto.vehicles[vin=1hgcm82633a004352].year"].change == "fixed"
    assert (result.base.correct, result.base.missed, result.base.invented) == (1, 5, 1)
    assert (result.adapter.correct, result.adapter.wrong, result.adapter.missed) == (3, 1, 2)
    assert result.adapter.recall == pytest.approx(3 / 6)


def test_a_table_without_an_identifier_is_matched_by_position():
    gold = {"policy": {"alternate_policy_identifiers": [{"identifier_value": _env("A")}]}}
    flat = keyed_flatten(gold, gold)
    assert "policy.alternate_policy_identifiers[#1].identifier_value" in flat


def test_the_workbooks_hold_a_summary_and_every_field(tmp_path):
    adapter = {"policy": {"policy_number": _env("PA-1")}}
    result = compare(GOLD, {}, adapter, lob="personal_auto")
    path = write_comparison(tmp_path / "comparison.xlsx", result, document="doc-1", line="personal_auto")
    book = load_workbook(path)
    assert book.sheetnames == ["Summary", "Fields"]
    header = [cell.value for cell in book["Fields"][1]]
    assert header[:6] == ["Field", "Gold", "Base", "Base result", "Base + adapter", "Base + adapter result"]
    assert book["Fields"].max_row == len(result.rows) + 1
    overview = write_overview(tmp_path / "summary.xlsx", [("doc-1", "personal_auto", result)])
    sheet = load_workbook(overview)["Documents"]
    assert sheet.cell(2, 1).value == "doc-1" and sheet.cell(3, 1).value.startswith("ALL DOCUMENTS")


# --------------------------------------------------------------------------
# Both routes through the serving pipeline
# --------------------------------------------------------------------------


class TwoAnswers(EchoBackend):
    """Answers like the base without an adapter, like the trained model with one."""

    BASE = '{"policy": {"policy_number": {"raw": "PA-9", "parsed": "PA-9", "page_ref": [1]}}}'
    TRAINED = '{"policy": {"policy_number": {"raw": "PA-1", "parsed": "PA-1", "page_ref": [1]}}}'

    def generate(self, messages, config, adapter=None):
        self.response = self.TRAINED if adapter else self.BASE
        return super().generate(messages, config, adapter=adapter)


@pytest.fixture
def client() -> BlobClient:
    return BlobClient(backend=InMemoryBackend(), container="main", raw_container="raw")


def _adapter(tmp_path: Path) -> Path:
    folder = tmp_path / "run" / "checkpoint-10"
    folder.mkdir(parents=True)
    (folder / "adapter_config.json").write_text(json.dumps({"r": 64}), encoding="utf-8")
    return folder


def _imported(client, source_id="policy_0001", tenant="compare"):
    client.write_json(paths.ocr_meta("policy", source_id, tenant), {"page_count": 1})
    client.write_bytes(paths.processed_page("policy", source_id, 1, "png", tenant), b"\x89PNG fake")
    client.write_text(paths.processed_page("policy", source_id, 1, "md", tenant), "Policy PA-1")
    client.write_json(paths.label_metadata("policy", source_id, tenant), {"lob": "personal_auto"})
    client.write_json(paths.golden_label("policy", source_id, tenant), {"policy": {"policy_number": _env("PA-1")}})


def test_one_document_goes_through_both_routes_with_one_model_load(client, tmp_path):
    from testing.compare_models import compare_documents

    _imported(client)
    backend = TwoAnswers()
    adapter_model = load_model("base", client, backend_impl=backend, adapter=str(_adapter(tmp_path)))
    base_model = without_adapter(adapter_model)
    assert base_model.backend is adapter_model.backend and base_model.default_adapter is None

    pdf = tmp_path / "in.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    compared = compare_documents(base_model, adapter_model, client, [("doc-1", "policy_0001", pdf)],
                                 tmp_path / "out", tenant="compare", images_root=tmp_path / "pages")

    folder = tmp_path / "out" / "doc-1"
    for name in ("document.pdf", "gold.json", "base.json", "adapter.json", "comparison.xlsx"):
        assert (folder / name).is_file(), name
    assert json.loads((folder / "base.json").read_text())["policy"]["policy_number"]["raw"] == "PA-9"
    assert json.loads((folder / "adapter.json").read_text())["policy"]["policy_number"]["raw"] == "PA-1"
    assert "effective_date" in json.loads((folder / "adapter.json").read_text())["policy"]   # every key
    (_name, _line, comparison), = compared
    assert comparison.base.wrong == 1 and comparison.adapter.correct == 1
    assert (tmp_path / "out" / "summary.xlsx").is_file()
    assert {call["adapter"] is None for call in backend.calls} == {True, False}


# --------------------------------------------------------------------------
# Inputs
# --------------------------------------------------------------------------


def test_a_pdf_with_its_gold_becomes_a_document_folder(tmp_path):
    from testing.compare_models import resolve_inputs

    pdf, gold = tmp_path / "policy.pdf", tmp_path / "gold.json"
    pdf.write_bytes(b"%PDF-1.4")
    gold.write_text(json.dumps({"line_of_business": ["homeowners"], "policy": {}}), encoding="utf-8")
    (folder,) = resolve_inputs(pdf, tmp_path / "staging", gold=gold)
    assert {p.name for p in folder.iterdir()} == {"document.pdf", "golden.json", "metadata.json"}
    assert json.loads((folder / "metadata.json").read_text())["lob"] == "homeowners"


def test_a_pdf_without_gold_or_line_is_refused(tmp_path):
    from testing.compare_models import CompareError, resolve_inputs

    pdf, gold = tmp_path / "policy.pdf", tmp_path / "gold.json"
    pdf.write_bytes(b"%PDF-1.4")
    with pytest.raises(CompareError, match="--gold"):
        resolve_inputs(pdf, tmp_path / "s")
    gold.write_text("{}", encoding="utf-8")
    with pytest.raises(CompareError, match="--lob"):
        resolve_inputs(pdf, tmp_path / "s", gold=gold)


def test_a_folder_of_document_folders_is_several_documents(tmp_path):
    from testing.compare_models import resolve_inputs

    for name in ("a", "b"):
        (tmp_path / name).mkdir()
        (tmp_path / name / "document.pdf").write_bytes(b"%PDF")
    assert [p.name for p in resolve_inputs(tmp_path, tmp_path / "s")] == ["a", "b"]
    assert resolve_inputs(tmp_path / "a", tmp_path / "s") == [tmp_path / "a"]


# --------------------------------------------------------------------------
# Download to the laptop through Blob
# --------------------------------------------------------------------------


def _run_folder(tmp_path: Path) -> Path:
    out = tmp_path / "compare-20261005-120000"
    (out / "doc-1").mkdir(parents=True)
    for name in ("gold.json", "base.json", "adapter.json", "comparison.xlsx"):
        (out / "doc-1" / name).write_text("{}", encoding="utf-8")
    (out / "summary.xlsx").write_text("x", encoding="utf-8")
    (out / "_inputs" / "doc-1").mkdir(parents=True)
    (out / "_inputs" / "doc-1" / "document.pdf").write_bytes(b"%PDF")
    return out


def test_a_finished_run_is_uploaded_for_download_and_removed_from_the_pod(tmp_path):
    from testing.compare_models import export_run

    raw = BlobClient(backend=InMemoryBackend(), container="main", raw_container="raw", context="ingestion")
    out = _run_folder(tmp_path)
    prefix, files = export_run(out, "compare", client=raw)
    assert prefix == "exports/compare/comparisons/compare-20261005-120000" and files == 5
    assert f"{prefix}/doc-1/comparison.xlsx" in raw.list(prefix + "/")
    assert not any("_inputs" in key for key in raw.list(prefix + "/"))
    assert not out.exists()


def test_keep_on_pod_keeps_the_folder(tmp_path):
    from testing.compare_models import export_run

    raw = BlobClient(backend=InMemoryBackend(), container="main", raw_container="raw", context="ingestion")
    out = _run_folder(tmp_path)
    export_run(out, "compare", client=raw, keep_local=True)
    assert (out / "summary.xlsx").is_file()


def test_exports_live_in_the_raw_container_under_its_rules():
    from artifact_registry.blob_client import AccessDeniedError

    key = paths.export_dir("comparisons", "run-1", "compare")
    assert paths.requires_raw_container(key) and paths.tenant_of(key) == "compare"
    general = BlobClient(backend=InMemoryBackend(), container="main", raw_container="raw")
    with pytest.raises(AccessDeniedError):
        general.read_bytes(f"{key}/summary.xlsx")
    with pytest.raises(ValueError):
        paths.export_dir("comparisons", "../escape", "compare")


def test_the_download_command_points_at_the_export_and_the_laptop_folder():
    from testing.compare_models import download_command

    command = download_command("exports/compare/comparisons/run-1", container="raw-docs", account="fideonstore")
    assert command == ('azcopy copy "https://fideonstore.blob.core.windows.net/raw-docs/exports/compare/'
                       'comparisons/run-1?<SAS>" "D:\Fine-Tuning-reports\comparisons" --recursive')
