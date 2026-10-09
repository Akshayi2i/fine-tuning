"""Base model vs base + adapter vs gold, field by field, as an Excel workbook.

Each document's three JSONs are flattened to one row per field. Table rows are
matched by their identifier (a VIN, a form number, a coverage name - the same
identifiers evaluation matches rows on), not by position, so a vehicle the model
listed second is compared with the gold's vehicle with the same VIN:
``auto.vehicles[vin=1hgcm82633a004352].year``. A table with no identifier is
matched by position (``[#2]``).

Each model's value gets a result against the gold value:

* **correct** - the same value (``values_agree``, the comparison every accuracy
  metric uses: dates, amounts and names normalised, numbers by their declared type);
* **wrong** - a different value;
* **missed** - the gold has a value, the model left it empty;
* **invented** - the model wrote a value where the gold has none.

Gold is compared as training and scoring use it: narrowed to what the line's
schema can hold (``common.canonical.schema_label``). The system-filled fields
(file name, page count) are left out of both sides.

:func:`field_report` gives one answer the same results as JSON, for every field
its line's schema declares - those neither the answer nor the gold holds too -
so a reader sees, out of all of them, which came back right, wrong, empty or not
at all.
"""

from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from functools import cache
from pathlib import Path
from typing import Any

RESULTS = ("correct", "wrong", "missed", "invented")

#: Cell colours per result.
FILLS = {"correct": "C6EFCE", "wrong": "FFC7CE", "missed": "FFEB9C", "invented": "E4D2F5"}


@dataclass
class FieldRow:
    path: str
    gold: Any
    base: Any
    adapter: Any
    base_result: str | None
    adapter_result: str | None
    gold_pages: list[int] = field(default_factory=list)

    @property
    def change(self) -> str:
        """``fixed`` when the adapter got right what the base did not; ``broke`` the reverse."""
        if self.gold in (None, "", []) and self.base_result is None and self.adapter_result is None:
            return ""
        base_ok, adapter_ok = self.base_result == "correct", self.adapter_result == "correct"
        if adapter_ok and not base_ok:
            return "fixed"
        if base_ok and not adapter_ok:
            return "broke"
        return ""


@dataclass
class ModelScore:
    correct: int = 0
    wrong: int = 0
    missed: int = 0
    invented: int = 0

    def add(self, result: str | None) -> None:
        if result:
            setattr(self, result, getattr(self, result) + 1)

    @property
    def gold_values(self) -> int:
        return self.correct + self.wrong + self.missed

    @property
    def written(self) -> int:
        return self.correct + self.wrong + self.invented

    @property
    def precision(self) -> float | None:
        return self.correct / self.written if self.written else None

    @property
    def recall(self) -> float | None:
        return self.correct / self.gold_values if self.gold_values else None

    @property
    def f1(self) -> float | None:
        total = self.written + self.gold_values
        return 2 * self.correct / total if total else None


@dataclass
class Comparison:
    rows: list[FieldRow]
    base: ModelScore
    adapter: ModelScore
    gold_outside_schema: int = 0


# --------------------------------------------------------------------------
# Flattening
# --------------------------------------------------------------------------


def _empty(value: Any) -> bool:
    return value is None or value == "" or value == [] or value == {}


def _is_rows(node: Any) -> bool:
    from common.canonical import is_field_value

    return (isinstance(node, list) and bool(node)
            and all(isinstance(r, dict) and not is_field_value(r) for r in node))


def _key_field(rows: list[dict[str, Any]]) -> str | None:
    """The identifier a table's rows are matched on, from the gold's rows when it has any."""
    from common.canonical import values_view
    from evaluation.metrics.field_accuracy import ROW_IDENTIFIERS, _infer_key_fields

    keys = _infer_key_fields(values_view(rows))
    return keys[0] if keys and keys[0] in ROW_IDENTIFIERS else None


def _labelled(rows: list[dict[str, Any]], key: str | None) -> list[tuple[str, dict[str, Any]]]:
    from common.canonical import values_view
    from evaluation.metrics.field_accuracy import _key_text

    out, seen = [], {}
    for index, row in enumerate(rows, start=1):
        text = _key_text(values_view(row.get(key))) if key else ""
        label = f"{key}={text}" if text else f"#{index}"
        seen[label] = seen.get(label, 0) + 1
        out.append((label if seen[label] == 1 else f"{label}~{seen[label]}", row))
    return out


def _join(prefix: str, key: str) -> str:
    return f"{prefix}.{key}" if prefix else key


def _leaves(node: Any, gold: Any, path: str, index: str, field: str) -> Iterator[tuple[str, str, str, Any]]:
    """Each leaf as ``(path, index path, schema field, leaf)`` - :func:`keyed_flatten`'s walk.

    The path labels a row by its identifier; the index path numbers it as the
    answer does (``coverages[3].premium``, the form review flags name); the
    schema field writes every row as ``[]`` (``coverages[].premium``).
    """
    from common.canonical import is_field_value

    if is_field_value(node):
        yield path, index, field, node
    elif isinstance(node, dict):
        source = gold if isinstance(gold, dict) else {}
        for key, value in node.items():
            yield from _leaves(value, source.get(key), _join(path, key), _join(index, key), _join(field, key))
    elif _is_rows(node):
        gold_rows = gold if _is_rows(gold) else []
        key = _key_field(gold_rows or node)
        gold_by_label = dict(_labelled(gold_rows, key))
        for position, (label, row) in enumerate(_labelled(node, key)):
            yield from _leaves(row, gold_by_label.get(label), f"{path}[{label}]", f"{index}[{position}]",
                               f"{field}[]")
    else:
        yield path, index, field, node


def keyed_flatten(node: Any, gold: Any = None, prefix: str = "",
                  out: dict[str, Any] | None = None) -> dict[str, Any]:
    """``{path: leaf}`` with table rows labelled by their identifier.

    A leaf is a field envelope, a list of envelopes (a set such as
    ``line_of_business``) or a bare value. ``gold`` is the gold label's node at
    the same place: its rows decide which field identifies a table, so the gold
    and both models label the same row the same way.
    """
    out = {} if out is None else out
    for path, _index, _field, leaf in _leaves(node, gold, prefix, prefix, prefix):
        out[path] = leaf
    return out


def _value(leaf: Any) -> Any:
    from common.canonical import values_view

    return values_view(leaf)


def _pages(leaf: Any) -> list[int]:
    from common.canonical import is_field_value

    leaves = leaf if isinstance(leaf, list) else [leaf]
    pages = {int(p) for item in leaves if is_field_value(item) for p in (item.get("page_ref") or [])}
    return sorted(pages)


def _result(gold: Any, got: Any, path: str) -> str | None:
    from evaluation.metrics.field_accuracy import values_agree

    if _empty(gold):
        return "invented" if not _empty(got) else None
    if _empty(got):
        return "missed"
    return "correct" if values_agree(gold, got, path) else "wrong"


def _stated(node: Any) -> int:
    from common.canonical import is_field_value

    if is_field_value(node):
        return int(node.get("raw") is not None or node.get("parsed") is not None)
    if isinstance(node, dict):
        return sum(_stated(v) for v in node.values())
    if isinstance(node, list):
        return sum(_stated(v) for v in node)
    return 0


def compare(gold_label: dict[str, Any] | None, base: dict[str, Any], adapter: dict[str, Any],
            *, doc_type: str = "policy", acord_form: str | None = None, lob: Any = None) -> Comparison:
    """One row per field any of the three holds a value for, with each model's result."""
    from common.canonical import schema_label, without_system_fields
    from common.label_mapping import map_label
    from common.structural_ids import comparable_for

    gold = without_system_fields(schema_label(gold_label or {}, doc_type, acord_form, lob))
    outside = _stated(without_system_fields(map_label(gold_label or {}, lob))) - _stated(gold)
    base, adapter = without_system_fields(base or {}), without_system_fields(adapter or {})
    # A common-model line's ids are each writer's own numbering: compared by
    # what its links name, never by id (common.structural_ids.comparable_view).
    gold, base, adapter = (comparable_for(doc, doc_type, acord_form, lob) for doc in (gold, base, adapter))

    flat_gold: dict[str, Any] = {}
    flat_base: dict[str, Any] = {}
    flat_adapter: dict[str, Any] = {}
    field_of: dict[str, str] = {}
    for flat, doc in ((flat_gold, gold), (flat_base, base), (flat_adapter, adapter)):
        for path, _index, name, leaf in _leaves(doc, gold, "", "", ""):
            flat[path], field_of[path] = leaf, name
    # A field the model is never asked for (what the pipeline fills, what another
    # path reads) is no model's to get right or wrong.
    unasked = _field_sets(doc_type, acord_form, lob)[1]

    rows: list[FieldRow] = []
    scores = (ModelScore(), ModelScore())
    for path in dict.fromkeys([*flat_gold, *flat_adapter, *flat_base]):
        if _within(field_of[path], unasked):
            continue
        gold_value = _value(flat_gold.get(path))
        base_value, adapter_value = _value(flat_base.get(path)), _value(flat_adapter.get(path))
        if _empty(gold_value) and _empty(base_value) and _empty(adapter_value):
            continue
        base_result, adapter_result = _result(gold_value, base_value, path), _result(gold_value, adapter_value, path)
        scores[0].add(base_result)
        scores[1].add(adapter_result)
        rows.append(FieldRow(path, gold_value, base_value, adapter_value, base_result, adapter_result,
                             _pages(flat_gold.get(path))))
    return Comparison(rows, scores[0], scores[1], max(outside, 0))


# --------------------------------------------------------------------------
# Every field of the schema, one answer
# --------------------------------------------------------------------------


#: What an answer holds for a field (:func:`field_report`).
FIELD_STATUSES: dict[str, str] = {
    "extracted": "the model wrote a value",
    "null": "the key is in the answer, with no value",
    "no_rows": "a column of a table the answer has no rows in",
    "not_in_output": "not in the answer: a row it does not have, or a key the schema does not allow to be null",
    "system": "never asked of the model: the pipeline fills it (a page count, a form's page range) "
              "or another path reads it (text_sections); not graded",
    "structural": "a row id: the writer's own numbering, compared through what names it; not graded",
}

#: Schema fields that hold no value from a page, with why.
_FIELD_NOTES = {"text_sections": "the full printed text, a tier of its own; not part of the extraction"}


def schema_fields(doc_type: str, acord_form: str | None = None, lob: Any = None, *,
                  asked: bool = False) -> list[str]:
    """Every field a schema declares, in schema order, named as :func:`keyed_flatten`
    names it with each row as ``[]``: ``carrier.name``, ``coverages[].limits[].amount``.

    A value in its envelope is one field, and so is a list of values. Any document
    type: a canonical schema refers only to its own ``$defs`` and is read as
    :func:`common.canonical.with_all_keys` reads the schema it fills; the flat
    schemas in ``schemas/`` (ACORD forms, Loss Runs) refer to other files and are
    read with those references inlined. ``asked`` gives the fields of the schema
    as the model is shown it (``common.schemas.resolved_schema``): without what
    the pipeline fills (a page count, a form's page range) or another path reads
    (``text_sections``) - the fields training teaches and scoring grades.
    """
    from common.schemas import _strip_prefixed, is_canonical, load_schema, resolved_schema

    schema = (_strip_prefixed(load_schema(doc_type, acord_form, lob), ("fideon:",))
              if is_canonical(doc_type, acord_form, lob) and not asked
              else resolved_schema(doc_type, acord_form, lob))
    defs = schema.get("$defs") or {}

    def resolve(sub: Any) -> dict[str, Any]:
        for _ in range(20):
            if not isinstance(sub, dict):
                return {}
            if "$ref" in sub:
                ref = sub["$ref"]
                sub = defs.get(ref.rsplit("/", 1)[-1]) if ref.startswith("#/$defs/") else None
                continue
            branches = sub.get("anyOf") or sub.get("oneOf")
            if branches and "properties" not in sub and "items" not in sub:
                sub = next((b for b in map(resolve, branches) if "properties" in b or "items" in b), {})
                continue
            return sub
        return {}

    def is_envelope(sub: dict[str, Any]) -> bool:
        return {"raw", "parsed", "page_ref"} <= set(sub.get("properties") or {})

    fields: list[str] = []

    def walk(sub: Any, path: str, depth: int) -> None:
        sub = resolve(sub)
        rows = resolve(sub.get("items")) if sub.get("type") == "array" or "items" in sub else {}
        if depth > 30 or is_envelope(sub):
            fields.append(path)
        elif sub.get("properties"):
            for key, child in sub["properties"].items():
                walk(child, _join(path, key), depth + 1)
        elif rows.get("properties") and not is_envelope(rows):
            walk(rows, f"{path}[]", depth + 1)
        elif path:
            fields.append(path)

    walk(schema, "", 0)
    return list(dict.fromkeys(fields))


def _field_sets(doc_type: str, acord_form: str | None, lob: Any) -> tuple[tuple[str, ...], frozenset[str]]:
    """``(every field the schema declares, those never asked of the model)``: the
    system fields, and what the schema as the model sees it leaves out. For a
    document no schema selects, no fields and only the system ones."""
    return _field_sets_for(doc_type, acord_form, tuple(lob) if isinstance(lob, list) else lob)


@cache
def _field_sets_for(doc_type: str, acord_form: str | None, lob: Any) -> tuple[tuple[str, ...], frozenset[str]]:
    from common.canonical import SYSTEM_SUPPLIED_FIELDS
    from common.schemas import SchemaError

    system = {f"{section}.{name}" for section, name in SYSTEM_SUPPLIED_FIELDS}
    line = list(lob) if isinstance(lob, tuple) else lob
    try:
        declared = schema_fields(doc_type, acord_form, line)
        asked = set(schema_fields(doc_type, acord_form, line, asked=True))
    except SchemaError:
        return (), frozenset(system)
    return tuple(declared), frozenset(system | (set(declared) - asked))


def _within(name: str, fields: Iterable[str]) -> bool:
    """Whether ``name`` is one of ``fields`` or inside one: a key of an open object
    such as ``text_sections``, whose leaves the schema does not name."""
    return any(name == f or name.startswith(f + ".") or name.startswith(f + "[") for f in fields)


def _section(field_name: str) -> str:
    return re.split(r"[.\[]", field_name, maxsplit=1)[0]


def field_report(answer: dict[str, Any] | None, gold: dict[str, Any] | None = None, *,
                 doc_type: str = "policy", acord_form: str | None = None, lob: Any = None,
                 review_flags: Iterable[str] = ()) -> dict[str, Any]:
    """Every field of the answer's schema, each with what the answer holds for it.

    ``fields`` lists every leaf of the answer and of the gold, rows matched by
    identifier as :func:`compare` matches them, then every field of the schema
    that neither holds - a column of a table with no rows, a key the schema does
    not allow to be null - so no field the line declares is left out. Each entry
    has its ``value`` and ``status`` (:data:`FIELD_STATUSES`); a value the model
    wrote has its ``confidence`` and ``page_ref``, and a field the pipeline
    flagged its review ``flags``. Against a gold label, an entry also has the
    ``gold`` value and the ``result``: correct, wrong, missed, invented, or empty
    - neither has a value, the field is not on this document. A table cell names
    its schema ``field`` as well.
    """
    from common.canonical import SYSTEM_SUPPLIED_FIELDS, is_field_value, schema_label, without_system_fields
    from common.schema_sections import structural_ids
    from common.schemas import SchemaError, schema_key
    from common.structural_ids import comparable_for

    answer = answer if isinstance(answer, dict) else {}
    graded = gold is not None
    declared, unasked = _field_sets(doc_type, acord_form, lob)
    system = {f"{section}.{name}" for section, name in SYSTEM_SUPPLIED_FIELDS}
    ids = {spec["field"] for spec in structural_ids(lob).values()}

    def not_graded(name: str, value: Any, status: str = "system") -> dict[str, Any]:
        entry = {"value": value, "status": status}
        note = next((text for field_name, text in _FIELD_NOTES.items() if _within(name, (field_name,))), None)
        if note:
            entry["note"] = note
        return entry

    # As compare() reads both sides: the gold narrowed to the schema, the system
    # fields out of scoring, ids compared through what they name.
    gold_view = (comparable_for(without_system_fields(schema_label(gold, doc_type, acord_form, lob)),
                                doc_type, acord_form, lob) if graded else {})
    answer_view = comparable_for(without_system_fields(answer), doc_type, acord_form, lob)
    keys = gold_view if graded else answer_view

    answered = {path: (index, name, leaf) for path, index, name, leaf in _leaves(answer_view, keys, "", "", "")}
    expected = {path: (name, leaf) for path, _index, name, leaf in _leaves(gold_view, gold_view, "", "", "")}
    by_index = {index: path for path, (index, _name, _leaf) in answered.items()}
    field_flags: dict[str, set[str]] = {}
    document_flags: list[str] = []
    for flag in review_flags:
        where, _, reason = str(flag).rpartition(":")
        if where in by_index:
            field_flags.setdefault(by_index[where], set()).add(reason)
        else:
            document_flags.append(str(flag))

    entries: dict[str, dict[str, Any]] = {}
    field_of: dict[str, str] = {}

    def add(path: str, name: str, entry: dict[str, Any]) -> None:
        if name != path:
            entry["field"] = name
        entries[path], field_of[path] = entry, name

    for path in dict.fromkeys([*answered, *expected]):
        _index, name, leaf = answered.get(path) or (None, expected[path][0], None)
        value = _value(leaf)
        if _within(name, unasked):
            add(path, name, not_graded(name, value))
            continue
        entry: dict[str, Any] = {
            "value": value,
            "status": "extracted" if not _empty(value) else "null" if path in answered else "not_in_output",
        }
        if is_field_value(leaf) and not _empty(value):
            confidence = leaf.get("confidence")
            entry["confidence"] = confidence.get("score") if isinstance(confidence, dict) else confidence
            entry["page_ref"] = list(leaf.get("page_ref") or [])
        if path in field_flags:
            entry["flags"] = sorted(field_flags[path])
        if graded:
            entry["gold"] = _value((expected.get(path) or (None, None))[1])
            entry["result"] = _result(entry["gold"], value, path) or "empty"
        add(path, name, entry)

    # What scoring leaves out, from the answer as served.
    for path, _index, name, leaf in _leaves(answer, keys, "", "", ""):
        last = name.rsplit(".", 1)[-1]
        if path in entries:
            continue
        if name in system:
            add(path, name, not_graded(name, _value(leaf)))
        elif last in ids or last == "part":
            add(path, name, not_graded(name, _value(leaf), "structural"))

    listed = set(field_of.values())
    for name in declared:
        if any(_within(other, (name,)) for other in listed):
            continue
        if name in unasked:
            add(name, name, not_graded(name, None))
            continue
        table = name[: name.rindex("[]") + 2] if "[]" in name else None
        entry = {"value": None,
                 "status": "no_rows" if table and not any(f.startswith(table + ".") for f in listed)
                 else "not_in_output"}
        if graded:
            entry["gold"], entry["result"] = None, "empty"
        add(name, name, entry)

    rank: dict[str, int] = {}
    for position, name in enumerate(declared):
        rank.setdefault(_section(name), position)
    ordered = sorted(entries, key=lambda path: rank.get(_section(field_of[path]), len(declared)))
    summary: dict[str, Any] = {
        "schema_fields": len(declared),
        "schema_fields_with_a_value": len({field_of[p] for p in entries if entries[p]["status"] == "extracted"}),
        "fields_listed": len(entries),
        "by_status": dict(Counter(entries[p]["status"] for p in ordered)),
    }
    if graded:
        tally = Counter(entries[p].get("result") for p in ordered)
        summary["by_result"] = {name: tally[name] for name in (*RESULTS, "empty") if tally[name]}
    try:
        schema = schema_key(doc_type, acord_form, lob)
    except SchemaError:   # no schema selects it: the answer's and the gold's fields only
        schema = None
    return {
        "schema": schema,
        "graded_against_gold": graded,
        "summary": summary,
        "document_flags": sorted(document_flags),
        "fields": {path: entries[path] for path in ordered},
    }


# --------------------------------------------------------------------------
# Workbooks
# --------------------------------------------------------------------------


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(str(v) for v in value)
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def _percent(value: float | None) -> Any:
    return round(value, 4) if value is not None else ""


def _style_header(sheet, row: int = 1) -> None:
    from openpyxl.styles import Font, PatternFill

    for cell in sheet[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")


def write_comparison(path: Path, comparison: Comparison, *, document: str, line: str,
                     base_label: str = "Base", adapter_label: str = "Base + adapter") -> Path:
    """``comparison.xlsx``: a Summary sheet and a Fields sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["Document", document])
    summary.append(["Line of business", line])
    summary.append(["Gold values compared", comparison.base.gold_values])
    summary.append(["Gold values outside the schema (not compared)", comparison.gold_outside_schema])
    summary.append([])
    header_row = summary.max_row + 1
    summary.append(["Measure", base_label, adapter_label, "Change"])
    _style_header(summary, header_row)
    for name in ("correct", "wrong", "missed", "invented"):
        a, b = getattr(comparison.base, name), getattr(comparison.adapter, name)
        summary.append([name.capitalize(), a, b, b - a])
    for name in ("precision", "recall", "f1"):
        a, b = getattr(comparison.base, name), getattr(comparison.adapter, name)
        change = round(b - a, 4) if a is not None and b is not None else ""
        summary.append([name.upper() if name == "f1" else name.capitalize(), _percent(a), _percent(b), change])
        for cell in summary[summary.max_row][1:4]:
            cell.number_format = "0.0%"
    fixed = sum(r.change == "fixed" for r in comparison.rows)
    broke = sum(r.change == "broke" for r in comparison.rows)
    summary.append(["Fields the adapter fixed / broke", fixed, broke, ""])
    summary.append([])
    for line_text in (
        "Correct: same value as gold (dates, amounts and names normalised). Wrong: a different value.",
        "Missed: gold has a value, the model left it empty. Invented: a value where gold has none.",
        "Precision = correct / values written. Recall = correct / gold values. F1 combines both.",
        "Table rows are matched by identifier (VIN, form number, coverage name ...), not position.",
    ):
        summary.append([line_text])
        summary[summary.max_row][0].font = Font(italic=True, color="595959")
    summary.column_dimensions["A"].width = 46
    for column in "BCD":
        summary.column_dimensions[column].width = 18

    fields = book.create_sheet("Fields")
    fields.append(["Field", "Gold", base_label, f"{base_label} result", adapter_label,
                   f"{adapter_label} result", "Change", "Gold page"])
    _style_header(fields)
    for row in comparison.rows:
        fields.append([row.path, _cell(row.gold), _cell(row.base), row.base_result or "",
                       _cell(row.adapter), row.adapter_result or "", row.change, _cell(row.gold_pages)])
        for column in (4, 6):
            result = fields.cell(fields.max_row, column).value
            if result in FILLS:
                fields.cell(fields.max_row, column).fill = PatternFill("solid", fgColor=FILLS[result])
    for column, width in zip("ABCDEFGH", (60, 32, 32, 14, 32, 16, 9, 10), strict=True):
        fields.column_dimensions[column].width = width
    for cells in fields.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 3, 5))
    fields.freeze_panes = "B2"
    fields.auto_filter.ref = fields.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def write_overview(path: Path, documents: Iterable[tuple[str, str, Comparison]]) -> Path:
    """``summary.xlsx``: one row per document, and the pooled totals."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Documents"
    sheet.append(["Document", "Line", "Gold values", "Base F1", "Adapter F1", "F1 change",
                  "Base recall", "Adapter recall", "Base precision", "Adapter precision",
                  "Adapter fixed", "Adapter broke"])
    _style_header(sheet)
    pooled = (ModelScore(), ModelScore())
    for name, line, comparison in documents:
        for score, source in zip(pooled, (comparison.base, comparison.adapter), strict=True):
            for result in RESULTS:
                setattr(score, result, getattr(score, result) + getattr(source, result))
        base, adapter = comparison.base, comparison.adapter
        sheet.append([name, line, base.gold_values, _percent(base.f1), _percent(adapter.f1),
                      _percent(adapter.f1 - base.f1) if base.f1 is not None and adapter.f1 is not None else "",
                      _percent(base.recall), _percent(adapter.recall),
                      _percent(base.precision), _percent(adapter.precision),
                      sum(r.change == "fixed" for r in comparison.rows),
                      sum(r.change == "broke" for r in comparison.rows)])
    base, adapter = pooled
    sheet.append(["ALL DOCUMENTS (pooled)", "", base.gold_values, _percent(base.f1), _percent(adapter.f1),
                  _percent(adapter.f1 - base.f1) if base.f1 is not None and adapter.f1 is not None else "",
                  _percent(base.recall), _percent(adapter.recall),
                  _percent(base.precision), _percent(adapter.precision), "", ""])
    for cells in sheet.iter_rows(min_row=2, min_col=4, max_col=10):
        for cell in cells:
            cell.number_format = "0.0%"
    sheet.column_dimensions["A"].width = 48
    sheet.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
